from flask import Flask, request, send_file, render_template_string
from flask_cors import CORS
from PIL import Image
import pytesseract
import io
from openpyxl import Workbook
from openpyxl.styles import Font
import os

app = Flask(__name__)

CORS(app)

HTML_PAGE = """
<!DOCTYPE html>
<html>
<head>
<meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>OCR Studio</title>
<link rel="icon" type="image/x-icon" href="/favicon.ico">
<link rel="icon" type="image/png" sizes="32x32" href="/favicon.ico">
<style>
*{box-sizing:border-box;font-family:Inter,system-ui,sans-serif}
body{
  margin:0;
  min-height:100vh;
  display:flex;
  align-items:center;
  justify-content:center;
  padding:20px;
  background-image: url('/background.jpg');
  background-size: cover;
  background-position: center;
  background-attachment: fixed;
  background-color: #7a8cff;
}
.card{
  width:100%;
  max-width:680px;
  background: rgba(255,255,255,0.25);
  backdrop-filter: blur(20px) saturate(180%);
  -webkit-backdrop-filter: blur(20px) saturate(180%);
  border-radius:24px;
  padding:32px;
  box-shadow: 0 8px 32px rgba(31,38,135,.2), inset 0 0 0 1px rgba(255,255,255,.4);
  border:1px solid rgba(255,255,255,.5);
}
h2{margin:0 0 8px;text-align:center;color:#111;text-shadow:0 1px 0 rgba(255,255,255,.8);font-size:26px;letter-spacing:-.5px}
.sub{text-align:center;color:rgba(0,0,0,.65);font-size:13px;margin-bottom:20px;font-weight:500}
.drop{
  border:1.5px dashed rgba(255,255,255,.8);
  border-radius:16px;
  padding:40px;
  text-align:center;
  background: rgba(255,255,255,.35);
  cursor:pointer;
  transition:.25s;
  color:#111;
  font-weight:600;
}
.drop:hover{background: rgba(255,255,255,.55); border-color:#fff; transform: translateY(-1px);}
.drop small{font-weight:400; opacity:.7}
.preview{display:grid;grid-template-columns:repeat(auto-fill,minmax(110px,1fr));gap:12px;margin-top:18px}
.item{border-radius:12px;overflow:hidden;background:rgba(255,255,255,.85);position:relative;box-shadow:0 4px 12px rgba(0,0,0,.1);border:1px solid rgba(255,255,255,.6)}
.item img{width:100%;height:100px;object-fit:cover;display:block}
.item span{display:block;padding:6px;font-size:11px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;background:rgba(255,255,255,.9)}
.item button{position:absolute;top:5px;right:5px;background:rgba(0,0,0,.65);color:#fff;border:none;border-radius:50%;width:24px;height:24px;cursor:pointer;backdrop-filter:blur(4px)}
.btn{
  width:100%;padding:16px;background:rgba(17,17,17,.9);color:#fff;border:1px solid rgba(255,255,255,.2);
  border-radius:14px;margin-top:18px;cursor:pointer;font-weight:700;font-size:15px;display:none;
  backdrop-filter:blur(10px); box-shadow:0 8px 20px rgba(0,0,0,.2); transition:.2s
}
.btn.show{display:block}
.btn:hover{transform:translateY(-1px); background:#000; box-shadow:0 12px 28px rgba(0,0,0,.3)}
.btn:disabled{background:rgba(100,100,100,.6)}
.log{
  background:rgba(17,17,17,.85);color:#7CFF7C;padding:12px;border-radius:10px;margin-top:14px;
  font-family:monospace;font-size:12px;white-space:pre-wrap; border:1px solid rgba(255,255,255,.15);
  backdrop-filter:blur(8px)
}
</style>
</head>
<body>
<div class="card">
<h2>OCR Studio</h2>
<div class="sub">File name in Row 1 • Preview then Download</div>
<div class="drop" id="dropZone" onclick="document.getElementById('images').click()">
<strong>✨ Click to select images</strong><br><small>or drag & drop here — multiple allowed</small>
<input type="file" id="images" multiple hidden accept="image/*">
</div>
<div id="preview" class="preview"></div>
<button class="btn" id="processBtn" onclick="upload()">Process & Download Excel</button>
<div id="log" class="log">Ready — Select images</div>
</div>
<script>
const input=document.getElementById('images'), preview=document.getElementById('preview'), btn=document.getElementById('processBtn'), log=document.getElementById('log'), dropZone=document.getElementById('dropZone');
let selectedFiles=[];
input.addEventListener('change', e=>{ handleFiles(e.target.files); });
dropZone.addEventListener('dragover', e=>{ e.preventDefault(); dropZone.style.background='rgba(255,255,255,.65)'; });
dropZone.addEventListener('dragleave', ()=>{ dropZone.style.background='rgba(255,255,255,.35)'; });
dropZone.addEventListener('drop', e=>{ e.preventDefault(); dropZone.style.background='rgba(255,255,255,.35)'; handleFiles(e.dataTransfer.files); });
function handleFiles(files){ selectedFiles=selectedFiles.concat(Array.from(files)); renderPreview(); }
function renderPreview(){
 preview.innerHTML='';
 if(selectedFiles.length===0){ btn.classList.remove('show'); log.textContent='Ready — Select images'; return; }
 selectedFiles.forEach((file,i)=>{
   const div=document.createElement('div'); div.className='item';
   const img=document.createElement('img'); img.src=URL.createObjectURL(file);
   const name=document.createElement('span'); name.textContent=file.name;
   const rm=document.createElement('button'); rm.textContent='×'; rm.onclick=()=>{ selectedFiles.splice(i,1); renderPreview(); };
   div.appendChild(img); div.appendChild(name); div.appendChild(rm); preview.appendChild(div);
 });
 btn.classList.add('show'); btn.textContent=`Process ${selectedFiles.length} Image(s) & Download`;
 log.textContent=`${selectedFiles.length} image(s) ready — See preview above`;
}
async function upload(){
 if(selectedFiles.length===0) return; btn.disabled=true; btn.textContent='Processing... ✨';
 log.textContent='OCR running...';
 const fd=new FormData(); selectedFiles.forEach(f=>fd.append('images', f));
 try{
  const res=await fetch('/ocr',{method:'POST',body:fd});
  if(!res.ok) throw new Error('Server error');
  const blob=await res.blob(); const url=URL.createObjectURL(blob);
  const a=document.createElement('a'); a.href=url; a.download='extracted_sheets.xlsx';
  document.body.appendChild(a); a.click(); a.remove();
  log.textContent='Success! Excel downloaded — Check your Downloads';
  btn.textContent='✅ Downloaded! Process again?';
 }catch(e){ log.textContent='Error: '+e.message; btn.textContent='Try Again'; }
 finally{ btn.disabled=false; }
}
</script>
</body>
</html>
"""

@app.route('/')
def home():
    return render_template_string(HTML_PAGE)

@app.route('/favicon.ico')
def favicon():
    if os.path.exists("favicon.ico"):
        return send_file("favicon.ico", mimetype="image/x-icon")
    return "", 404

@app.route('/background.jpg')
def background():
    if os.path.exists("background.jpg"):
        return send_file("background.jpg", mimetype="image/jpeg")
    return "", 404

@app.route('/ocr', methods=['POST'])
def ocr():
    files = request.files.getlist('images')
    wb = Workbook()
    wb.remove(wb.active)
    for f in files:
        raw_name = f.filename.rsplit('.', 1)[0]
        safe_sheet = "".join(c if c.isalnum() or c in (" ", "_") else "_" for c in raw_name)[:31] or "Sheet1"
        ws = wb.create_sheet(title=safe_sheet)
        try:
            img = Image.open(f.stream).convert("L")
            text = pytesseract.image_to_string(img, lang='eng', config='--oem 3 --psm 6')
        except Exception as e:
            text = f"OCR Error: {str(e)}"
        ws['A1'] = f.filename
        ws['A1'].font = Font(bold=True, size=12, color="0000FF")
        row_num = 3
        for line in text.splitlines():
            cleaned = line.strip()
            if cleaned == "":
                row_num += 1
                continue
            ws.cell(row=row_num, column=1, value=cleaned)
            row_num += 1
        ws.column_dimensions['A'].width = 80
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="extracted_sheets.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

if __name__ == '__main__':
    port = int(os.environ.get("PORT", 5003))
    app.run(host='0.0.0.0', port=port)
